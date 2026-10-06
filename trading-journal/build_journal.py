"""Builds NIFTY_Options_Trading_Journal.xlsx (formulas only, no macros).

Run:  python3 build_journal.py  -> writes the workbook next to this script.
"""
import datetime as dt
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

OUT = Path(__file__).with_name("NIFTY_Options_Trading_Journal.xlsx")

# ---------------------------------------------------------------- styles
FONT = "Arial"
BLUE, BLACK, WHITE = "0000FF", "000000", "FFFFFF"
NAVY, GREY, LIGHT = "1F3864", "595959", "F2F2F2"
INPUT_FILL = PatternFill("solid", fgColor="FFF9E5")   # pale yellow = type here
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

F_IN = Font(name=FONT, size=10, color=BLUE)
F_CALC = Font(name=FONT, size=10, color=BLACK)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color=GREY)
F_HDR = Font(name=FONT, size=10, bold=True, color=WHITE)
F_SEC = Font(name=FONT, size=11, bold=True, color=NAVY)

HDR_IN = PatternFill("solid", fgColor="2F5597")    # header of an input column
HDR_CALC = PatternFill("solid", fgColor="404040")  # header of a formula column
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")

INR = '"₹"#,##0;[Red]-"₹"#,##0;"-"'
INR2 = '"₹"#,##0.00;[Red]-"₹"#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
PCT2 = '0.000%'
PX = '0.00'
NUM = '#,##0'
DATE = 'dd-mmm-yy'
TIME = 'hh:mm'
RFMT = '0.00"R";[Red]-0.00"R";"0R"'

wb = Workbook()


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB


def section(ws, cell, text, span=1):
    c = ws[cell]
    c.value = text
    c.font = F_SEC
    c.fill = SEC_FILL
    if span > 1:
        col = c.column
        for i in range(1, span):
            ws.cell(row=c.row, column=col + i).fill = SEC_FILL


def put(ws, cell, value, fmt=None, inp=False, bold=False):
    c = ws[cell]
    c.value = value
    c.font = Font(name=FONT, size=10, bold=bold, color=BLUE if inp else BLACK)
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = INPUT_FILL
    c.border = BOX
    return c


def add_table(ws, ref, tname, style="TableStyleLight1"):
    t = Table(displayName=tname, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


def list_dv(ws, rng, list_name, prompt=None):
    dv = DataValidation(type="list", formula1=f"={list_name}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Pick from list",
                        error="Choose a value from the dropdown (edit choices on LISTS).")
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Input", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


# ======================================================================
# LISTS (built first so named ranges exist for everything else)
# ======================================================================
wsL = wb.active
wsL.title = "LISTS"
LISTS = {
    "List_Instrument": ("Instrument", ["NIFTY", "SENSEX", "BANKNIFTY", "OTHER"]),
    "List_OptType": ("CE/PE", ["CE", "PE"]),
    "List_Side": ("Buy/Sell", ["Sell", "Buy"]),
    "List_Strategy": ("Strategy", [
        "Naked Short Call", "Naked Short Put", "Short Strangle", "Short Straddle",
        "Bear Call Spread", "Bull Put Spread", "Iron Condor", "Iron Fly",
        "Hedge / Long Option", "Other"]),
    "List_ExitReason": ("Exit Reason", ["Target", "SL", "Time exit", "Panic",
                                        "Revenge", "Rule break", "EOD"]),
    "List_RuleBroken": ("Rule Broken", [
        "None", "R1 Oversized position", "R2 No hard SL placed",
        "R3 Moved / ignored SL", "R4 Held short past expiry cutoff",
        "R5 Revenge trade", "R6 Exceeded max trades", "R7 Traded after daily limit",
        "R8 Gave back profit (no trail)", "R9 Unplanned / FOMO entry",
        "R10 Traded in poor state"]),
    "List_Emotion": ("Emotion", ["Calm", "Focused", "Confident", "Anxious", "Fearful",
                                 "Greedy", "FOMO", "Frustrated", "Angry / Revenge",
                                 "Bored", "Euphoric", "Tired"]),
    "List_Grade": ("Setup Grade", ["A", "B", "C"]),
    "List_YN": ("Y/N", ["Y", "N"]),
    "List_Bias": ("Market Bias", ["Bullish", "Mildly Bullish", "Range-bound",
                                  "Mildly Bearish", "Bearish", "No view"]),
    "List_Scale": ("Scale 1-5", [1, 2, 3, 4, 5]),
    "List_Weekday": ("Weekday", ["Mon", "Tue", "Wed", "Thu", "Fri"]),
}
title(wsL, "LISTS – dropdown sources",
      "Edit the values below to change dropdown choices. Keep 'None' in Rule Broken. "
      "Each list is an Excel Table and a named range (List_*).")
col = 1
for nm, (hdr, vals) in LISTS.items():
    L = get_column_letter(col)
    wsL[f"{L}4"] = hdr
    for i, v in enumerate(vals):
        put(wsL, f"{L}{5 + i}", v, inp=True)
    last = 4 + len(vals)
    name(nm, f"LISTS!${L}$5:${L}${last}")
    add_table(wsL, f"{L}4:{L}{last}", "tbl" + nm.replace("List_", "L"))
    wsL.column_dimensions[L].width = max(14, max(len(str(v)) for v in vals) + 3)
    col += 1

# entry-time buckets (two columns: start time + label)
Ls, Ll = get_column_letter(col), get_column_letter(col + 1)
wsL[f"{Ls}4"], wsL[f"{Ll}4"] = "Bucket Start", "Entry-Time Bucket"
buckets = [(dt.time(0, 0), "Pre-open (<09:15)"), (dt.time(9, 15), "09:15-10:00"),
           (dt.time(10, 0), "10:00-11:30"), (dt.time(11, 30), "11:30-13:30"),
           (dt.time(13, 30), "13:30-14:45"), (dt.time(14, 45), "14:45-15:30 (last 45m)")]
for i, (t, lab) in enumerate(buckets):
    put(wsL, f"{Ls}{5 + i}", t, TIME, inp=True)
    put(wsL, f"{Ll}{5 + i}", lab, inp=True)
blast = 4 + len(buckets)
name("List_BucketStart", f"LISTS!${Ls}$5:${Ls}${blast}")
name("List_BucketLabel", f"LISTS!${Ll}$5:${Ll}${blast}")
add_table(wsL, f"{Ls}4:{Ll}{blast}", "tblLBuckets")
wsL.column_dimensions[Ls].width = 13
wsL.column_dimensions[Ll].width = 24
wsL.freeze_panes = "A5"

# ======================================================================
# RULES
# ======================================================================
wsR = wb.create_sheet("RULES", 0)
title(wsR, "NIFTY Weekly Options – Trading Rules & Journal",
      "Blue text on yellow = inputs you edit. Black = formulas (do not overwrite).")
wsR.column_dimensions["A"].width = 3
wsR.column_dimensions["B"].width = 46
wsR.column_dimensions["C"].width = 16
wsR.column_dimensions["D"].width = 62
wsR.column_dimensions["E"].width = 3

section(wsR, "B4", "HOW TO USE (target: < 5 minutes per trade)", 3)
howto = [
    "1. Once: set your inputs below (capital, risk %, limits, lot size, charges). Everything else recalculates.",
    "2. Before 09:15: fill one row on PRE-MARKET. If it says NO TRADE / HALF SIZE, respect it.",
    "3. Before each entry: use the Position-Size Calculator below. Lots must not exceed 'Lots allowed'.",
    "4. After each exit: fill one row on TRADE LOG (blue columns only). Grey-header columns are automatic.",
    "5. Check DASHBOARD after every trade. If it shows STOP TRADING TODAY – close the terminal.",
    "6. Friday/weekend: complete WEEKLY REVIEW (20 min). Pick ONE change for next week.",
    "7. Sample trades on TRADE LOG / PRE-MARKET are examples – delete their blue inputs when you start.",
    "8. For spreads/condors you may log the net credit as one row and set Legs = 2 or 4 (charges approximate).",
]
for i, t in enumerate(howto):
    c = wsR[f"B{5 + i}"]
    c.value = t
    c.font = F_CALC
    wsR.merge_cells(f"B{5 + i}:D{5 + i}")

section(wsR, "B14", "INPUTS – risk limits", 3)
put(wsR, "C14", "Value", bold=True)
put(wsR, "D14", "Notes / source", bold=True)
risk_inputs = [
    (15, "Trading capital (₹)", 270000, INR, "Capital", "Current account value. Update monthly (not daily)."),
    (16, "Risk per trade (% of capital)", 0.01, PCT, "RiskPct", "Default 1%. Max 2% for small accounts."),
    (17, "Daily loss limit (% of capital)", 0.02, PCT, "DailyLossPct", "Hit it = stop for the day. No exceptions."),
    (18, "Weekly max drawdown (% of capital)", 0.06, PCT, "WeeklyDDPct", "Week-to-date net loss. Hit it = no trading until next Monday."),
    (19, "Expiry-day exit cutoff (time)", dt.time(14, 45), TIME, "ExpiryCutoff", "All short options flat by this time on expiry day (DTE 0)."),
    (20, "Max lots per position", 2, NUM, "MaxLots", "Hard cap even if the risk calc allows more."),
    (21, "Max trades per day", 3, NUM, "MaxTradesDay", "Counts every logged row (each leg if you log legs)."),
    (22, "Max consecutive losses → stop for day", 2, NUM, "MaxConsecLosses", "Anti-tilt circuit breaker."),
    (23, "NIFTY lot size (units)", 65, NUM, "LotSize", "NSE: 65 from Jan-2026 series (was 75). Re-check NSE circulars each Jun/Dec."),
    (24, "Weekly expiry day", "Tuesday", None, "ExpiryDay", "NIFTY weekly expiry = Tuesday since Sep-2025 (SEBI expiry-day framework)."),
]
for r, lab, val, fmt, nm, note in risk_inputs:
    put(wsR, f"B{r}", lab)
    put(wsR, f"C{r}", val, fmt, inp=True)
    put(wsR, f"D{r}", note)
    name(nm, f"RULES!$C${r}")

section(wsR, "B26", "INPUTS – Zerodha F&O option charges (verify on zerodha.com/charges)", 3)
charge_inputs = [
    (27, "Brokerage per executed order (₹)", 20, INR, "BrokeragePerOrder", "Flat ₹20/order for F&O options."),
    (28, "STT – sell side, on premium", 0.0015, PCT2, "STTRate", "0.15% from 1-Apr-2026 (Union Budget 2026-27; was 0.10%)."),
    (29, "NSE transaction charge, on premium turnover", 0.0003503, '0.00000%', "ExchRate", "0.03503% (NSE options)."),
    (30, "SEBI turnover fee (₹10 per crore)", 0.000001, '0.000000%', "SEBIRate", "₹10 / ₹1,00,00,000 = 0.0001%."),
    (31, "Stamp duty – buy side", 0.00003, PCT2, "StampRate", "0.003% on buy value."),
    (32, "GST on brokerage + exchange + SEBI fees", 0.18, PCT, "GSTRate", "18%."),
]
for r, lab, val, fmt, nm, note in charge_inputs:
    put(wsR, f"B{r}", lab)
    put(wsR, f"C{r}", val, fmt, inp=True)
    put(wsR, f"D{r}", note)
    name(nm, f"RULES!$C${r}")

section(wsR, "B34", "AUTO-CALCULATED LIMITS (₹)", 3)
calc_rows = [
    (35, "Max risk per trade (1R) ₹", "=Capital*RiskPct", INR, "RiskPerTrade", "Your 1R. Planned SL loss must be ≤ this."),
    (36, "Daily loss limit ₹", "=Capital*DailyLossPct", INR, "DailyLossLimit", "Your ₹1.2 lakh day would be ~22x this limit."),
    (37, "Weekly drawdown limit ₹", "=Capital*WeeklyDDPct", INR, "WeeklyDDLimit", ""),
    (38, "Max SL distance for 1 lot (premium pts)", "=RiskPerTrade/LotSize", PX, "MaxSLPts1Lot", "If your SL is wider than this, the trade is too big even at 1 lot."),
    (39, "Losing days in a row to hit weekly limit", "=IFERROR(WeeklyDDLimit/DailyLossLimit,0)", '0.0', None, "Weekly limit = this many max-loss days."),
]
for r, lab, f, fmt, nm, note in calc_rows:
    put(wsR, f"B{r}", lab)
    put(wsR, f"C{r}", f, fmt)
    put(wsR, f"D{r}", note)
    if nm:
        name(nm, f"RULES!$C${r}")

section(wsR, "B41", "POSITION-SIZE CALCULATOR (use before every entry)", 3)
put(wsR, "B42", "Entry premium (₹)")
put(wsR, "C42", 51, PX, inp=True)
put(wsR, "D42", "Option price you will sell at.")
put(wsR, "B43", "Hard stop-loss premium (₹)")
put(wsR, "C43", 66, PX, inp=True)
put(wsR, "D43", "Place this as an SL-M / SL-L order immediately after entry.")
put(wsR, "B44", "SL distance (pts)", )
put(wsR, "C44", "=ABS(C43-C42)", PX)
put(wsR, "B45", "Risk per lot ₹")
put(wsR, "C45", "=C44*LotSize", INR)
put(wsR, "B46", "Lots allowed (risk-based, capped at Max lots)", bold=True)
put(wsR, "C46", "=IF(C45<=0,0,MIN(MaxLots,INT(RiskPerTrade/C45)))", NUM)
put(wsR, "D46", '=IF(C46=0,"SL too wide for your 1R – skip or tighten / use a spread.","Max lots for this trade.")')
put(wsR, "B47", "Total risk at allowed lots ₹")
put(wsR, "C47", "=C46*C45", INR)

section(wsR, "B49", "MY RULES (built from the inputs above)", 3)
rules = [
    '="R1  Size: risk ≤ "&TEXT(RiskPerTrade,"₹#,##0")&" per trade (1R) and never more than "&MaxLots&" lots."',
    '="R2  Hard SL: a stop-loss ORDER goes in within 30 seconds of entry. Mental stops do not count."',
    '="R3  Never move or cancel an SL further away. Only trail it toward profit."',
    '="R4  Expiry day (DTE 0): all shorts flat by "&TEXT(ExpiryCutoff,"hh:mm")&". No new shorts after 14:00."',
    '="R5  After any SL hit: 15-minute break before the next entry. No re-entry in the same strike."',
    '="R6  Max "&MaxTradesDay&" trades per day."',
    '="R7  Day net ≤ -"&TEXT(DailyLossLimit,"₹#,##0")&" or "&MaxConsecLosses&" losses in a row = done for the day. Week ≤ -"&TEXT(WeeklyDDLimit,"₹#,##0")&" = done for the week."',
    '="R8  Profit protection: at 40% of premium captured, SL to cost; at 60%, trail SL to lock half the open profit."',
    '="R9  Only A/B setups written on PRE-MARKET. No FOMO entries."',
    '="R10 Emotional state ≤ 2, sleep < 6h or \'trading to recover\' = Y → no trading or half size."',
]
for i, f in enumerate(rules):
    c = wsR[f"B{50 + i}"]
    c.value = f
    c.font = F_CALC
    wsR.merge_cells(f"B{50 + i}:D{50 + i}")

section(wsR, "B61", "COLOUR LEGEND", 3)
put(wsR, "B62", "Blue text, yellow fill", inp=True)
put(wsR, "D62", "Input – type here")
put(wsR, "B63", "Black text")
put(wsR, "D63", "Formula – do not overwrite")
put(wsR, "B64", "Blue column header (logs)").fill = HDR_IN
wsR["B64"].font = F_HDR
put(wsR, "D64", "Input column on TRADE LOG / PRE-MARKET")
put(wsR, "B65", "Dark-grey column header (logs)").fill = HDR_CALC
wsR["B65"].font = F_HDR
put(wsR, "D65", "Automatic column")
put(wsR, "B66", "Red fill").fill = RED_FILL
put(wsR, "D66", "Rule / limit breach")
wsR.sheet_view.showGridLines = False

# ======================================================================
# TRADE LOG
# ======================================================================
wsT = wb.create_sheet("TRADE LOG", 1)
FIRST, LAST = 6, 305  # 300 pre-built rows

# (key, header, kind, width, fmt, group, formula template)
# kind: "in" = input, "f" = formula. In templates {Key} -> that column's cell on the same row.
TL = [
    ("No", "Trade #", "f", 7, "0", "BASICS", '=IF({Date}="","",COUNT($B${F}:{Date}))'),
    ("Date", "Date", "in", 11, DATE, "BASICS", None),
    ("Day", "Day", "f", 6, None, "BASICS", '=IF({Date}="","",TEXT({Date},"ddd"))'),
    ("Expiry", "Expiry Date", "in", 11, DATE, "BASICS", None),
    ("DTE", "DTE", "f", 6, "0", "BASICS", '=IF(OR({Date}="",{Expiry}=""),"",MAX(0,NETWORKDAYS({Date},{Expiry})-1))'),
    ("Instr", "Instrument", "in", 10, None, "BASICS", None),
    ("Strike", "Strike", "in", 9, "0", "BASICS", None),
    ("Type", "CE/PE", "in", 7, None, "BASICS", None),
    ("Side", "Buy/Sell", "in", 8, None, "BASICS", None),
    ("Strat", "Strategy", "in", 17, None, "BASICS", None),
    ("PosID", "Position ID", "in", 9, None, "BASICS", None),
    ("Legs", "Legs", "in", 6, "0", "BASICS", None),
    ("Lots", "Lots", "in", 6, "0", "BASICS", None),
    ("Qty", "Quantity", "f", 9, NUM, "BASICS", '=IF({Lots}="","",{Lots}*LotSize)'),
    ("ETime", "Entry Time", "in", 8, TIME, "ENTRY", None),
    ("EPx", "Entry Price", "in", 9, PX, "ENTRY", None),
    ("Spot", "NIFTY Spot at Entry", "in", 10, "#,##0.00", "ENTRY", None),
    ("Dist", "Distance from Spot (pts)", "f", 10, "#,##0;[Red]-#,##0", "ENTRY", '=IF(OR({Strike}="",{Spot}="",{Type}=""),"",IF({Type}="CE",{Strike}-{Spot},{Spot}-{Strike}))'),
    ("DistPct", "Distance %", "f", 9, '0.00%;[Red]-0.00%', "ENTRY", '=IF({Dist}="","",{Dist}/{Spot})'),
    ("IV", "IV at Entry %", "in", 8, "0.0", "ENTRY", None),
    ("Delta", "Delta at Entry", "in", 8, "0.00", "ENTRY", None),
    ("Margin", "Margin Used ₹", "in", 11, INR, "ENTRY", None),
    ("Defined", "Max Loss Defined? (Y/N)", "in", 9, None, "ENTRY", None),
    ("SL", "Stop-Loss Price", "in", 9, PX, "RISK", None),
    ("Tgt", "Target Price", "in", 9, PX, "RISK", None),
    ("Risk", "Planned Risk ₹ (1R)", "f", 11, INR, "RISK", '=IF(OR({SL}="",{EPx}="",{Qty}=""),"",ABS({SL}-{EPx})*{Qty})'),
    ("RR", "Planned Reward:Risk", "f", 9, '0.00', "RISK", '=IF(OR({Risk}="",{Tgt}=""),"",IF({Risk}=0,"",ABS({EPx}-{Tgt})*{Qty}/{Risk}))'),
    ("RiskPctC", "Risk % of Capital", "f", 9, '0.00%', "RISK", '=IF({Risk}="","",{Risk}/Capital)'),
    ("SizeChk", "Size Check", "f", 11, None, "RISK", '=IF({Date}="","",IF({SL}="","NO SL",IF(OR(N({Lots})>MaxLots,N({RiskPctC})>RiskPct),"OVERSIZED","OK")))'),
    ("XTime", "Exit Time", "in", 8, TIME, "EXIT", None),
    ("XPx", "Exit Price", "in", 9, PX, "EXIT", None),
    ("XReason", "Exit Reason", "in", 11, None, "EXIT", None),
    ("MAEPx", "MAE Price (worst seen)", "in", 9, PX, "PATH", None),
    ("MFEPx", "MFE Price (best seen)", "in", 9, PX, "PATH", None),
    ("MAE", "MAE ₹", "f", 10, INR, "PATH", '=IF(OR({MAEPx}="",{EPx}="",{Qty}=""),"",ABS({MAEPx}-{EPx})*{Qty})'),
    ("MFE", "MFE ₹", "f", 10, INR, "PATH", '=IF(OR({MFEPx}="",{EPx}="",{Qty}=""),"",ABS({EPx}-{MFEPx})*{Qty})'),
    ("Gross", "Gross P&L ₹", "f", 11, INR, "RESULTS", '=IF(OR({XPx}="",{EPx}="",{Qty}=""),"",IF({Side}="Buy",{XPx}-{EPx},{EPx}-{XPx})*{Qty})'),
    ("GiveBack", "Profit Given Back ₹ (MFE - Gross)", "f", 11, INR, "RESULTS", '=IF(OR({MFE}="",{Gross}=""),"",MAX(0,{MFE}-{Gross}))'),
    ("Turn", "Premium Turnover ₹", "f", 11, INR, "CHARGES", '=IF({Gross}="","",({EPx}+{XPx})*{Qty})'),
    ("Brok", "Brokerage ₹", "f", 9, INR2, "CHARGES", '=IF({Gross}="","",BrokeragePerOrder*2*MAX(1,N({Legs})))'),
    ("STT", "STT ₹", "f", 9, INR2, "CHARGES", '=IF({Gross}="","",STTRate*IF({Side}="Buy",{XPx},{EPx})*{Qty})'),
    ("Exch", "Exchange Txn ₹", "f", 9, INR2, "CHARGES", '=IF({Gross}="","",ExchRate*{Turn})'),
    ("SEBI", "SEBI Fee ₹", "f", 8, INR2, "CHARGES", '=IF({Gross}="","",SEBIRate*{Turn})'),
    ("Stamp", "Stamp Duty ₹", "f", 8, INR2, "CHARGES", '=IF({Gross}="","",StampRate*IF({Side}="Buy",{EPx},{XPx})*{Qty})'),
    ("GST", "GST ₹", "f", 8, INR2, "CHARGES", '=IF({Gross}="","",GSTRate*({Brok}+{Exch}+{SEBI}))'),
    ("Chg", "Total Charges ₹", "f", 10, INR2, "RESULTS", '=IF({Gross}="","",{Brok}+{STT}+{Exch}+{SEBI}+{Stamp}+{GST})'),
    ("Net", "Net P&L ₹", "f", 11, INR, "RESULTS", '=IF({Gross}="","",{Gross}-{Chg})'),
    ("R", "R-Multiple", "f", 8, RFMT, "RESULTS", '=IF({Net}="","",{Net}/IF(N({Risk})>0,{Risk},RiskPerTrade))'),
    ("PctCap", "% of Capital", "f", 8, '0.00%;[Red]-0.00%', "RESULTS", '=IF({Net}="","",{Net}/Capital)'),
    ("Followed", "Followed Rules? (Y/N)", "in", 9, None, "BEHAVIOUR", None),
    ("Broken", "Rule Broken", "in", 24, None, "BEHAVIOUR", None),
    ("EmoB", "Emotion Before", "in", 12, None, "BEHAVIOUR", None),
    ("EmoA", "Emotion After", "in", 12, None, "BEHAVIOUR", None),
    ("Grade", "Setup Grade (A/B/C)", "in", 8, None, "BEHAVIOUR", None),
    ("Shot", "Screenshot Link", "in", 14, None, "BEHAVIOUR", None),
    ("Lesson", "Lesson", "in", 44, None, "BEHAVIOUR", None),
    ("WL", "Win/Loss", "f", 7, None, "AUTO", '=IF({Net}="","",IF({Net}>0,"Win",IF({Net}<0,"Loss","BE")))'),
    ("Bucket", "Entry-Time Bucket", "f", 16, None, "AUTO", '=IF({ETime}="","",LOOKUP({ETime},List_BucketStart,List_BucketLabel))'),
    ("Week", "Week Start (Mon)", "f", 10, DATE, "AUTO", '=IF({Date}="","",{Date}-WEEKDAY({Date},3))'),
    ("NoDay", "Trade # of Day", "f", 7, "0", "AUTO", '=IF({Date}="","",COUNTIF($B${F}:{Date},{Date}))'),
    ("DayBefore", "Day P&L Before This Trade ₹", "f", 11, INR, "AUTO", '=IF({Date}="","",SUMIFS(${NetL}${F}:{Net},$B${F}:{Date},{Date})-N({Net}))'),
    ("ExpBreach", "Expiry Cutoff Check", "f", 10, None, "AUTO", '=IF(OR({Date}="",{DTE}=""),"",IF(AND({DTE}=0,{Side}="Sell",MAX(N({ETime}),N({XTime}))>ExpiryCutoff),"BREACH","OK"))'),
    ("AutoFlag", "Auto Rule Check", "f", 10, None, "AUTO", '=IF({Date}="","",IF(OR({SizeChk}<>"OK",{ExpBreach}="BREACH",{NoDay}>MaxTradesDay,{DayBefore}<=-DailyLossLimit,OR({XReason}="Panic",{XReason}="Revenge",{XReason}="Rule break")),"CHECK","OK"))'),
    ("Cum", "Cumulative Net ₹", "f", 11, INR, "AUTO", '=IF({Net}="","",SUM(${NetL}${F}:{Net}))'),
    ("Eq", "Equity ₹", "f", 11, INR, "AUTO", '=IF({Cum}="","",Capital+{Cum})'),
    ("Peak", "Peak Equity ₹", "f", 11, INR, "AUTO", '=IF({Eq}="","",MAX(Capital,MAX(${EqL}${F}:{Eq})))'),
    ("DD", "Drawdown ₹", "f", 10, INR, "AUTO", '=IF({Eq}="","",{Peak}-{Eq})'),
    ("DDPct", "Drawdown %", "f", 9, PCT, "AUTO", '=IF({DD}="","",{DD}/{Peak})'),
    ("Streak", "Loss Streak", "f", 7, "0", "AUTO", '=IF({Net}="","",IF({Net}<0,N({StreakPrev})+1,0))'),
]
COL = {k: get_column_letter(i + 1) for i, (k, *_rest) in enumerate(TL)}
assert COL["Date"] == "B"


def render(tpl, r):
    out = tpl.replace("{F}", str(FIRST)).replace("{NetL}", COL["Net"]).replace("{EqL}", COL["Eq"])
    out = out.replace("{StreakPrev}", f'{COL["Streak"]}{r - 1}')
    return re.sub(r"\{(\w+)\}", lambda m: f"{COL[m.group(1)]}{r}", out)


title(wsT, "TRADE LOG – one row per trade (or per net-credit spread)",
      "Fill BLUE-header columns only. Dark-grey headers are automatic. 300 rows pre-built; "
      "rows 6-10 are samples (incl. the 51 → 30 → 70 expiry-day short call).")
wsT["A3"] = ("MAE = worst price seen against you; MFE = best price seen in your favour (for a short: "
             "MAE is the HIGHEST premium, MFE the LOWEST). R-multiple uses Planned Risk; if no SL was "
             "set it uses your rule 1R (RULES).")
wsT["A3"].font = F_SUB

# group band row 4
runs = []  # contiguous (group, first col, last col)
for i, (k, h, kind, w, fmt, grp, tpl) in enumerate(TL, start=1):
    if runs and runs[-1][0] == grp:
        runs[-1][2] = i
    else:
        runs.append([grp, i, i])
    c = wsT.cell(row=5, column=i, value=h)
    c.font = F_HDR
    c.fill = HDR_IN if kind == "in" else HDR_CALC
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = BOX
    wsT.column_dimensions[get_column_letter(i)].width = w
GROUP_FILLS = ["DDEBF7", "E2EFDA", "FCE4D6", "FFF2CC", "EDEDED", "D9E1F2", "F8CBAD", "E4DFEC", "D0CECE"]
for gi, (grp, a, b) in enumerate(runs):
    c = wsT.cell(row=4, column=a, value=grp)
    c.font = F_BOLD
    c.alignment = Alignment(horizontal="center")
    for j in range(a, b + 1):
        wsT.cell(row=4, column=j).fill = PatternFill("solid", fgColor=GROUP_FILLS[gi % len(GROUP_FILLS)])
    if b > a:
        wsT.merge_cells(start_row=4, start_column=a, end_row=4, end_column=b)
wsT.row_dimensions[5].height = 42

T = dt.time
samples = [
    # Date, Expiry, Instr, Strike, Type, Side, Strat, PosID, Legs, Lots, ETime, EPx, Spot, IV, Delta, Margin, Defined,
    # SL, Tgt, XTime, XPx, XReason, MAEPx, MFEPx, Followed, Broken, EmoB, EmoA, Grade, Shot, Lesson
    dict(Date=dt.date(2026, 9, 24), Expiry=dt.date(2026, 9, 29), Instr="NIFTY", Strike=26000, Type="CE", Side="Sell",
         Strat="Bear Call Spread", PosID="P001", Legs=2, Lots=2, ETime=T(9, 48), EPx=24.5, Spot=25812, IV=11.8,
         Delta=0.16, Margin=62000, Defined="Y", SL=36, Tgt=8, XTime=T(13, 10), XPx=8.2, XReason="Target",
         MAEPx=27, MFEPx=7.9, Followed="Y", Broken="None", EmoB="Calm", EmoA="Confident", Grade="A", Shot="",
         Lesson="26000/26200 CE spread, net credit. Plan executed; booked 67% of credit."),
    dict(Date=dt.date(2026, 9, 25), Expiry=dt.date(2026, 9, 29), Instr="NIFTY", Strike=25650, Type="PE", Side="Sell",
         Strat="Naked Short Put", PosID="P002", Legs=1, Lots=1, ETime=T(10, 5), EPx=38, Spot=25790, IV=12.4,
         Delta=0.22, Margin=198000, Defined="N", SL=57, Tgt=15, XTime=T(11, 20), XPx=57.5, XReason="SL",
         MAEPx=58, MFEPx=33, Followed="Y", Broken="None", EmoB="Calm", EmoA="Frustrated", Grade="B", Shot="",
         Lesson="Valid setup, SL order filled. A good loss = -1R."),
    dict(Date=dt.date(2026, 9, 25), Expiry=dt.date(2026, 9, 29), Instr="NIFTY", Strike=25600, Type="PE", Side="Sell",
         Strat="Naked Short Put", PosID="P003", Legs=1, Lots=1, ETime=T(11, 32), EPx=44, Spot=25735, IV=13.1,
         Delta=0.27, Margin=201000, Defined="N", SL=None, Tgt=10, XTime=T(14, 55), XPx=112, XReason="Revenge",
         MAEPx=118, MFEPx=41, Followed="N", Broken="R5 Revenge trade", EmoB="Angry / Revenge", EmoA="Frustrated",
         Grade="C", Shot="",
         Lesson="Re-entered 12 min after SL to 'get it back'. No SL. Lost 3.4x the first loss."),
    dict(Date=dt.date(2026, 9, 28), Expiry=dt.date(2026, 9, 29), Instr="NIFTY", Strike=25500, Type="PE", Side="Sell",
         Strat="Bull Put Spread", PosID="P004", Legs=2, Lots=2, ETime=T(10, 20), EPx=18, Spot=25705, IV=12.0,
         Delta=0.14, Margin=58000, Defined="Y", SL=27, Tgt=6, XTime=T(14, 30), XPx=7.5, XReason="Time exit",
         MAEPx=21.5, MFEPx=6.8, Followed="Y", Broken="None", EmoB="Focused", EmoA="Calm", Grade="A", Shot="",
         Lesson="25500/25300 PE spread. Time exit before close as per plan."),
    dict(Date=dt.date(2026, 9, 29), Expiry=dt.date(2026, 9, 29), Instr="NIFTY", Strike=25850, Type="CE", Side="Sell",
         Strat="Bear Call Spread", PosID="P005", Legs=2, Lots=5, ETime=T(10, 40), EPx=51, Spot=25815, IV=14.2,
         Delta=0.38, Margin=182000, Defined="Y", SL=66, Tgt=15, XTime=T(15, 8), XPx=70, XReason="Panic",
         MAEPx=74, MFEPx=30, Followed="N", Broken="R4 Held short past expiry cutoff", EmoB="Greedy", EmoA="Fearful",
         Grade="B", Shot="",
         Lesson="Expiry day: 51 → 30 (41% captured) → 70. No trail, SL only mental, 5 lots (cap 2), "
                "held past 14:45 into gamma spike. Booking/trailing near 30-40 = +₹3,400 to +₹6,800 instead of -₹6,175."),
]
for i in range(FIRST, LAST + 1):
    s = samples[i - FIRST] if i - FIRST < len(samples) else {}
    for j, (k, h, kind, w, fmt, grp, tpl) in enumerate(TL, start=1):
        c = wsT.cell(row=i, column=j)
        if kind == "f":
            c.value = render(tpl, i)
            c.font = F_CALC
        else:
            v = s.get(k)
            c.value = v if v not in ("",) else None
            c.font = F_IN
        if fmt:
            c.number_format = fmt
TLREF = f"A5:{COL['Streak']}{LAST}"
add_table(wsT, TLREF, "tblTrades", "TableStyleLight15")
wsT.freeze_panes = "C6"
# group (outline) the charge breakdown so it can be collapsed
for k in ("Turn", "Brok", "STT", "Exch", "SEBI", "Stamp", "GST"):
    wsT.column_dimensions[COL[k]].outlineLevel = 1

# named ranges for every trade-log column
TLN = {}
for k, h, *_ in TL:
    nm = "TL_" + k
    TLN[k] = nm
    name(nm, f"'TRADE LOG'!${COL[k]}${FIRST}:${COL[k]}${LAST}")


def rng(k):
    return f"{COL[k]}{FIRST}:{COL[k]}{LAST}"


# data validation
list_dv(wsT, rng("Instr"), "List_Instrument")
list_dv(wsT, rng("Type"), "List_OptType")
list_dv(wsT, rng("Side"), "List_Side")
list_dv(wsT, rng("Strat"), "List_Strategy", "For spreads, log net credit and set Legs.")
list_dv(wsT, rng("Defined"), "List_YN")
list_dv(wsT, rng("XReason"), "List_ExitReason")
list_dv(wsT, rng("Followed"), "List_YN")
list_dv(wsT, rng("Broken"), "List_RuleBroken")
list_dv(wsT, rng("EmoB"), "List_Emotion")
list_dv(wsT, rng("EmoA"), "List_Emotion")
list_dv(wsT, rng("Grade"), "List_Grade")
for k, lo, hi, msg in (("Lots", 1, 100, "Whole number of lots."), ("Legs", 1, 4, "1 = single option, 2 = spread, 4 = condor/fly.")):
    dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True, error=msg, showInputMessage=True, promptTitle=k, prompt=msg)
    wsT.add_data_validation(dv)
    dv.add(rng(k))
for k in ("Date", "Expiry"):
    dv = DataValidation(type="date", operator="greaterThan", formula1="43831", allow_blank=True,
                        showErrorMessage=True, error="Enter a date (e.g. 29-09-2026).")
    wsT.add_data_validation(dv)
    dv.add(rng(k))
for k in ("ETime", "XTime"):
    dv = DataValidation(type="time", operator="between", formula1="0.375", formula2="0.6667", allow_blank=True,
                        showErrorMessage=True, error="Enter a market-hours time, e.g. 10:40.")
    wsT.add_data_validation(dv)
    dv.add(rng(k))
for k in ("EPx", "XPx", "SL", "Tgt", "MAEPx", "MFEPx"):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, error="Premium must be ≥ 0.")
    wsT.add_data_validation(dv)
    dv.add(rng(k))

# conditional formatting
B = f"$B{FIRST}"
for k in ("Net", "Gross", "R", "PctCap"):
    wsT.conditional_formatting.add(rng(k), CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    wsT.conditional_formatting.add(rng(k), CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL))
wsT.conditional_formatting.add(rng("SizeChk"), FormulaRule(
    formula=[f'AND({COL["SizeChk"]}{FIRST}<>"",{COL["SizeChk"]}{FIRST}<>"OK")'], fill=RED_FILL, font=Font(bold=True, color="9C0006")))
wsT.conditional_formatting.add(rng("ExpBreach"), FormulaRule(
    formula=[f'{COL["ExpBreach"]}{FIRST}="BREACH"'], fill=RED_FILL, font=Font(bold=True, color="9C0006")))
wsT.conditional_formatting.add(rng("AutoFlag"), FormulaRule(
    formula=[f'{COL["AutoFlag"]}{FIRST}="CHECK"'], fill=AMBER_FILL, font=Font(bold=True)))
wsT.conditional_formatting.add(rng("XReason"), FormulaRule(
    formula=[f'OR({COL["XReason"]}{FIRST}="Panic",{COL["XReason"]}{FIRST}="Revenge",{COL["XReason"]}{FIRST}="Rule break")'],
    fill=RED_FILL))
wsT.conditional_formatting.add(rng("Followed"), FormulaRule(formula=[f'{COL["Followed"]}{FIRST}="N"'], fill=RED_FILL))
wsT.conditional_formatting.add(rng("GiveBack"), FormulaRule(
    formula=[f'AND(ISNUMBER({COL["GiveBack"]}{FIRST}),{COL["GiveBack"]}{FIRST}>=RiskPerTrade)'], fill=AMBER_FILL))
wsT.conditional_formatting.add(rng("NoDay"), FormulaRule(
    formula=[f'AND(ISNUMBER({COL["NoDay"]}{FIRST}),{COL["NoDay"]}{FIRST}>MaxTradesDay)'], fill=RED_FILL))
wsT[f'{COL["Lesson"]}5'].comment = Comment("One sentence: what will you do differently next time?", "Journal")
wsT[f'{COL["MAEPx"]}5'].comment = Comment(
    "Short option: the HIGHEST premium while you held it. Long option: the LOWEST.", "Journal")
wsT[f'{COL["MFEPx"]}5'].comment = Comment(
    "Short option: the LOWEST premium while you held it. Long option: the HIGHEST.", "Journal")
wsT[f'{COL["DTE"]}5'].comment = Comment(
    "Trading days to expiry (Mon-Fri; exchange holidays not excluded). 0 = expiry day.", "Journal")
wsT[f'{COL["Chg"]}5'].comment = Comment(
    "Zerodha F&O options: ₹20/order ×2 per leg + STT 0.15% sell side + NSE 0.03503% + SEBI ₹10/cr + "
    "stamp 0.003% buy side + 18% GST. Rates on RULES. For net-credit spread rows, STT/exchange are "
    "approximate (computed on net premium).", "Journal")

# ======================================================================
# PRE-MARKET
# ======================================================================
wsP = wb.create_sheet("PRE-MARKET", 1)
PF, PL = 6, 105
PM = [
    ("Date", "Date", "in", 11, DATE, None),
    ("Day", "Day", "f", 6, None, '=IF({Date}="","",TEXT({Date},"ddd"))'),
    ("Expiry", "Next Expiry", "in", 11, DATE, None),
    ("DTE", "DTE", "f", 6, "0", '=IF(OR({Date}="",{Expiry}=""),"",MAX(0,NETWORKDAYS({Date},{Expiry})-1))'),
    ("PH", "Prev High", "in", 10, "#,##0.00", None),
    ("PLo", "Prev Low", "in", 10, "#,##0.00", None),
    ("PC", "NIFTY Prev Close", "in", 10, "#,##0.00", None),
    ("Open", "Open / Pre-open", "in", 10, "#,##0.00", None),
    ("Gap", "Gap (pts)", "f", 8, "#,##0;[Red]-#,##0", '=IF(OR({Open}="",{PC}=""),"",{Open}-{PC})'),
    ("GapPct", "Gap %", "f", 8, '0.00%;[Red]-0.00%', '=IF({Gap}="","",{Gap}/{PC})'),
    ("VIX", "India VIX", "in", 8, "0.00", None),
    ("Move", "VIX 1-SD Day Move (pts)", "f", 10, "#,##0", '=IF(OR({VIX}="",{PC}=""),"",{PC}*{VIX}/100/SQRT(252))'),
    ("RLo", "Expected Low (Open - 1SD)", "f", 10, "#,##0", '=IF(OR({Move}="",{Open}=""),"",{Open}-{Move})'),
    ("RHi", "Expected High (Open + 1SD)", "f", 10, "#,##0", '=IF(OR({Move}="",{Open}=""),"",{Open}+{Move})'),
    ("Piv", "Pivot (P)", "f", 10, "#,##0.0", '=IF(OR({PH}="",{PLo}="",{PC}=""),"",({PH}+{PLo}+{PC})/3)'),
    ("BC", "CPR BC", "f", 10, "#,##0.0", '=IF({Piv}="","",({PH}+{PLo})/2)'),
    ("TC", "CPR TC", "f", 10, "#,##0.0", '=IF({Piv}="","",2*{Piv}-{BC})'),
    ("CPRW", "CPR Width %", "f", 8, "0.000%", '=IF({Piv}="","",ABS({TC}-{BC})/{Piv})'),
    ("R1", "R1", "f", 10, "#,##0.0", '=IF({Piv}="","",2*{Piv}-{PLo})'),
    ("S1", "S1", "f", 10, "#,##0.0", '=IF({Piv}="","",2*{Piv}-{PH})'),
    ("MaxPain", "Max Pain", "in", 9, "0", None),
    ("CallOI", "Highest Call OI Strike", "in", 9, "0", None),
    ("PutOI", "Highest Put OI Strike", "in", 9, "0", None),
    ("Bias", "Market Bias", "in", 13, None, None),
    ("Setups", "Planned Setups (strikes, entry trigger, SL)", "in", 42, None, None),
    ("Emo", "Emotional State (1-5)", "in", 8, "0", None),
    ("Sleep", "Sleep (hrs)", "in", 7, "0.0", None),
    ("Stress", "Stress (1-5)", "in", 7, "0", None),
    ("Recover", "Trading to Recover Losses? (Y/N)", "in", 10, None, None),
    ("PrevPnL", "Last Session Net P&L ₹", "f", 11, INR, '=IF({Date}="","",IFERROR(SUMIFS(TL_Net,TL_Date,_xlfn.MAXIFS(TL_Date,TL_Date,"<"&{Date})),0))'),
    ("Go", "GO / NO-GO", "f", 20, None, '=IF({Date}="","",IF(OR({Recover}="Y",AND({Emo}<>"",N({Emo})<=2),AND({Sleep}<>"",N({Sleep})<6),N({Stress})>=4),"NO TRADE / HALF SIZE","GO"))'),
]
PCOL = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(PM)}


def prender(tpl, r):
    return re.sub(r"\{(\w+)\}", lambda m: f"{PCOL[m.group(1)]}{r}", tpl)


title(wsP, "PRE-MARKET PLAN – fill before 09:15",
      "Score: Emotional State 1 = very poor … 5 = excellent; Stress 1 = none … 5 = extreme. "
      "VIX 1-SD move = Prev Close × VIX/100/√252 (≈68% of days stay inside).")
wsP["A3"] = "Levels: CPR (Pivot/BC/TC), R1/S1 from previous day H/L/C. Max pain and OI strikes from the option chain (Sensibull/Opstra/NSE)."
wsP["A3"].font = F_SUB
for i, (k, h, kind, w, fmt, tpl) in enumerate(PM, start=1):
    c = wsP.cell(row=5, column=i, value=h)
    c.font = F_HDR
    c.fill = HDR_IN if kind == "in" else HDR_CALC
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    wsP.column_dimensions[get_column_letter(i)].width = w
wsP.row_dimensions[5].height = 54
pm_samples = [
    dict(Date=dt.date(2026, 9, 24), Expiry=dt.date(2026, 9, 29), PH=25868, PLo=25742, PC=25790, Open=25805, VIX=11.6,
         MaxPain=25800, CallOI=26000, PutOI=25600, Bias="Range-bound",
         Setups="Sell 26000/26200 CE spread above 25,830 rejection. SL 36. 2 lots max.", Emo=4, Sleep=7, Stress=2, Recover="N"),
    dict(Date=dt.date(2026, 9, 25), Expiry=dt.date(2026, 9, 29), PH=25851, PLo=25760, PC=25795, Open=25801, VIX=12.2,
         MaxPain=25800, CallOI=26000, PutOI=25700, Bias="Mildly Bullish",
         Setups="Sell 25650 PE if 25,760 holds. SL 57. 1 lot naked only.", Emo=3, Sleep=6.5, Stress=3, Recover="N"),
    dict(Date=dt.date(2026, 9, 28), Expiry=dt.date(2026, 9, 29), PH=25812, PLo=25640, PC=25672, Open=25690, VIX=12.8,
         MaxPain=25700, CallOI=26000, PutOI=25500, Bias="Range-bound",
         Setups="Bull put spread 25500/25300 near S1. Time exit 14:30.", Emo=3, Sleep=6, Stress=3, Recover="Y"),
    dict(Date=dt.date(2026, 9, 29), Expiry=dt.date(2026, 9, 29), PH=25748, PLo=25661, PC=25731, Open=25790, VIX=13.9,
         MaxPain=25750, CallOI=25900, PutOI=25700, Bias="Mildly Bearish",
         Setups="EXPIRY DAY. Max 2 lots. Exit all shorts by 14:45. Sell 25850 CE spread only on rejection.",
         Emo=2, Sleep=5, Stress=4, Recover="Y"),
]
for r in range(PF, PL + 1):
    s = pm_samples[r - PF] if r - PF < len(pm_samples) else {}
    for j, (k, h, kind, w, fmt, tpl) in enumerate(PM, start=1):
        c = wsP.cell(row=r, column=j)
        if kind == "f":
            c.value = prender(tpl, r)
            c.font = F_CALC
        else:
            c.value = s.get(k)
            c.font = F_IN
        if fmt:
            c.number_format = fmt
add_table(wsP, f"A5:{PCOL['Go']}{PL}", "tblPremarket", "TableStyleLight15")
wsP.freeze_panes = "B6"


def prng(k):
    return f"{PCOL[k]}{PF}:{PCOL[k]}{PL}"


list_dv(wsP, prng("Bias"), "List_Bias")
list_dv(wsP, prng("Emo"), "List_Scale", "1 = very poor … 5 = excellent")
list_dv(wsP, prng("Stress"), "List_Scale", "1 = none … 5 = extreme")
list_dv(wsP, prng("Recover"), "List_YN", "Be honest. Y = NO TRADE / HALF SIZE.")
wsP.conditional_formatting.add(prng("Go"), FormulaRule(formula=[f'{PCOL["Go"]}{PF}="NO TRADE / HALF SIZE"'],
                                                       fill=RED_FILL, font=Font(bold=True, color="9C0006")))
wsP.conditional_formatting.add(prng("Go"), FormulaRule(formula=[f'{PCOL["Go"]}{PF}="GO"'], fill=GREEN_FILL))
wsP.conditional_formatting.add(prng("Recover"), FormulaRule(formula=[f'{PCOL["Recover"]}{PF}="Y"'], fill=RED_FILL))

# ======================================================================
# DAILY SUMMARY
# ======================================================================
wsD = wb.create_sheet("DAILY SUMMARY", 3)
DF, DL = 6, 205
DS = [
    ("Date", "Date", DATE, None),
    ("Day", "Day", None, '=IF({Date}="","",TEXT({Date},"ddd"))'),
    ("Trades", "Trades", "0", '=IF({Date}="","",COUNTIF(TL_Date,{Date}))'),
    ("Wins", "Wins", "0", '=IF({Date}="","",COUNTIFS(TL_Date,{Date},TL_Net,">0"))'),
    ("Losses", "Losses", "0", '=IF({Date}="","",COUNTIFS(TL_Date,{Date},TL_Net,"<0"))'),
    ("Gross", "Gross P&L ₹", INR, '=IF({Date}="","",SUMIFS(TL_Gross,TL_Date,{Date}))'),
    ("Chg", "Charges ₹", INR, '=IF({Date}="","",SUMIFS(TL_Chg,TL_Date,{Date}))'),
    ("Net", "Net P&L ₹", INR, '=IF({Date}="","",SUMIFS(TL_Net,TL_Date,{Date}))'),
    ("Cum", "Cumulative Net ₹", INR, '=IF({Date}="","",SUM($H${F}:{Net}))'),
    ("Eq", "Equity ₹", INR, '=IF({Date}="","",Capital+{Cum})'),
    ("Peak", "Peak Equity ₹", INR, '=IF({Date}="","",MAX(Capital,MAX($J${F}:{Eq})))'),
    ("DD", "Drawdown from Peak ₹", INR, '=IF({Date}="","",{Peak}-{Eq})'),
    ("DDPct", "Drawdown %", PCT, '=IF({Date}="","",{DD}/{Peak})'),
    ("Week", "Week Start (Mon)", DATE, '=IF({Date}="","",{Date}-WEEKDAY({Date},3))'),
    ("WTD", "Week-to-Date Net ₹", INR, '=IF({Date}="","",SUMIFS($H${F}:{Net},$N${F}:{Week},{Week}))'),
    ("Adh", "Rule Adherence %", PCT, '=IF({Date}="","",IFERROR(COUNTIFS(TL_Date,{Date},TL_Followed,"Y")/{Trades},0))'),
    ("Breaks", "Rule Breaks", "0", '=IF({Date}="","",COUNTIFS(TL_Date,{Date},TL_Followed,"N"))'),
    ("MaxLotsUsed", "Max Lots Used", "0", '=IF({Date}="","",_xlfn.MAXIFS(TL_Lots,TL_Date,{Date}))'),
    ("GiveBack", "Profit Given Back ₹", INR, '=IF({Date}="","",SUMIFS(TL_GiveBack,TL_Date,{Date}))'),
    ("ExpB", "Expiry Cutoff Breaches", "0", '=IF({Date}="","",COUNTIFS(TL_Date,{Date},TL_ExpBreach,"BREACH"))'),
    ("DayFlag", "Daily Loss Limit", None, '=IF({Date}="","",IF({Net}<=-DailyLossLimit,"BREACH","OK"))'),
    ("TrFlag", "Max Trades", None, '=IF({Date}="","",IF({Trades}>MaxTradesDay,"BREACH","OK"))'),
    ("WkFlag", "Weekly Drawdown", None, '=IF({Date}="","",IF({WTD}<=-WeeklyDDLimit,"BREACH","OK"))'),
    ("Status", "Day Status", None, '=IF({Date}="","",IF(OR({DayFlag}="BREACH",{TrFlag}="BREACH",{WkFlag}="BREACH",{ExpB}>0,{MaxLotsUsed}>MaxLots),"LIMIT BROKEN","OK"))'),
]
DCOL = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(DS)}
assert DCOL["Net"] == "H" and DCOL["Eq"] == "J" and DCOL["Week"] == "N"


def drender(tpl, r):
    out = tpl.replace("{F}", str(DF))
    return re.sub(r"\{(\w+)\}", lambda m: f"{DCOL[m.group(1)]}{r}", out)


title(wsD, "DAILY SUMMARY – fully automatic from TRADE LOG",
      "Dates are pulled from TRADE LOG in ascending order. Do not type here. Weekly drawdown = week-to-date net loss.")
for i, (k, h, fmt, tpl) in enumerate(DS, start=1):
    c = wsD.cell(row=5, column=i, value=h)
    c.font = F_HDR
    c.fill = HDR_CALC
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    wsD.column_dimensions[get_column_letter(i)].width = 11 if k not in ("Status",) else 15
wsD.row_dimensions[5].height = 42
for r in range(DF, DL + 1):
    for j, (k, h, fmt, tpl) in enumerate(DS, start=1):
        c = wsD.cell(row=r, column=j)
        if k == "Date":
            c.value = ('=IFERROR(SMALL(TL_Date,1),"")' if r == DF else
                       f'=IF(A{r - 1}="","",IFERROR(SMALL(TL_Date,COUNTIF(TL_Date,"<="&A{r - 1})+1),""))')
        else:
            c.value = drender(tpl, r)
        c.font = F_CALC
        if fmt:
            c.number_format = fmt
add_table(wsD, f"A5:{DCOL['Status']}{DL}", "tblDaily", "TableStyleLight15")
wsD.freeze_panes = "B6"
for k in DCOL:
    name("DS_" + k, f"'DAILY SUMMARY'!${DCOL[k]}${DF}:${DCOL[k]}${DL}")


def drng(k):
    return f"{DCOL[k]}{DF}:{DCOL[k]}{DL}"


for k in ("DayFlag", "TrFlag", "WkFlag"):
    wsD.conditional_formatting.add(drng(k), FormulaRule(formula=[f'{DCOL[k]}{DF}="BREACH"'], fill=RED_FILL,
                                                        font=Font(bold=True, color="9C0006")))
wsD.conditional_formatting.add(drng("Status"), FormulaRule(formula=[f'{DCOL["Status"]}{DF}="LIMIT BROKEN"'],
                                                           fill=RED_FILL, font=Font(bold=True, color="9C0006")))
wsD.conditional_formatting.add(drng("Status"), FormulaRule(formula=[f'{DCOL["Status"]}{DF}="OK"'], fill=GREEN_FILL))
for k in ("Net", "WTD"):
    wsD.conditional_formatting.add(drng(k), CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")))
wsD.conditional_formatting.add(drng("Adh"), FormulaRule(
    formula=[f'AND(ISNUMBER({DCOL["Adh"]}{DF}),{DCOL["Adh"]}{DF}<1)'], fill=AMBER_FILL))

# ======================================================================
# DASHBOARD
# ======================================================================
wsB = wb.create_sheet("DASHBOARD", 4)
wsB.sheet_view.showGridLines = False
title(wsB, "PERFORMANCE DASHBOARD", "Formulas only – updates as you log trades. 'Latest session' = most recent date on TRADE LOG.")
widths = {"A": 2, "B": 34, "C": 15, "D": 52, "E": 2, "F": 30, "G": 9, "H": 9, "I": 13, "J": 9, "K": 13, "L": 2}
for k, v in widths.items():
    wsB.column_dimensions[k].width = v

# --- STOP banner + today block
wsB.merge_cells("B4:K4")
put(wsB, "B4", '=IF(C14="STOP","STOP TRADING TODAY  –  "&C15,"OK TO TRADE  –  within all limits")', bold=True)
wsB["B4"].font = Font(name=FONT, size=16, bold=True)
wsB["B4"].alignment = Alignment(horizontal="center", vertical="center")
wsB.row_dimensions[4].height = 34
wsB.conditional_formatting.add("B4:K4", FormulaRule(formula=['$C$14="STOP"'], fill=PatternFill("solid", fgColor="C00000"),
                                                    font=Font(bold=True, color=WHITE)))
wsB.conditional_formatting.add("B4:K4", FormulaRule(formula=['$C$14<>"STOP"'], fill=GREEN_FILL,
                                                    font=Font(bold=True, color="006100")))

section(wsB, "B6", "LATEST SESSION – RISK STATUS", 3)
today = [
    (7, "Latest session date", '=IF(COUNT(TL_Date)=0,"",MAX(TL_Date))', DATE, "Most recent date on TRADE LOG."),
    (8, "Trades today / max", '=IF(C7="",0,COUNTIF(TL_Date,C7))', "0", '="Max "&MaxTradesDay&" per day"'),
    (9, "Net P&L today ₹", '=IF(C7="",0,SUMIFS(TL_Net,TL_Date,C7))', INR, '="Limit: -"&TEXT(DailyLossLimit,"₹#,##0")'),
    (10, "Daily loss limit remaining ₹", '=DailyLossLimit+MIN(0,C9)', INR, "≤ 0 → stop."),
    (11, "Week-to-date net ₹", '=IF(C7="",0,SUMIFS(TL_Net,TL_Week,C7-WEEKDAY(C7,3)))', INR, '="Weekly limit: -"&TEXT(WeeklyDDLimit,"₹#,##0")'),
    (12, "Current consecutive losses", '=IFERROR(INDEX(TL_Streak,MATCH(9.99E+307,TL_Streak)),0)', "0", '="Stop at "&MaxConsecLosses'),
    (13, "Expiry-cutoff breaches today", '=IF(C7="",0,COUNTIFS(TL_Date,C7,TL_ExpBreach,"BREACH"))', "0", '="Shorts flat by "&TEXT(ExpiryCutoff,"hh:mm")&" on expiry"'),
    (14, "STOP TRADING TODAY?", '=IF(OR(C9<=-DailyLossLimit,C8>=MaxTradesDay,C11<=-WeeklyDDLimit,C12>=MaxConsecLosses),"STOP","OK")', None, "Any one trigger = stop."),
    (15, "Reason", '=IF(C14="OK","-",TRIM(IF(C9<=-DailyLossLimit,"Daily loss limit hit. ","")&IF(C8>=MaxTradesDay,"Max trades reached. ","")&IF(C11<=-WeeklyDDLimit,"Weekly drawdown limit hit. ","")&IF(C12>=MaxConsecLosses,"Consecutive-loss limit hit. ","")))', None, ""),
]
for r, lab, f, fmt, note in today:
    put(wsB, f"B{r}", lab)
    put(wsB, f"C{r}", f, fmt)
    put(wsB, f"D{r}", note)
wsB.merge_cells("C15:D15")
wsB.conditional_formatting.add("C9", FormulaRule(formula=["C9<=-DailyLossLimit"], fill=RED_FILL))
wsB.conditional_formatting.add("C10", FormulaRule(formula=["C10<=0"], fill=RED_FILL))
wsB.conditional_formatting.add("C8", FormulaRule(formula=["C8>=MaxTradesDay"], fill=RED_FILL))
wsB.conditional_formatting.add("C11", FormulaRule(formula=["C11<=-WeeklyDDLimit"], fill=RED_FILL))
wsB.conditional_formatting.add("C12", FormulaRule(formula=["C12>=MaxConsecLosses"], fill=RED_FILL))
wsB.conditional_formatting.add("C13", FormulaRule(formula=["C13>0"], fill=RED_FILL))
wsB.conditional_formatting.add("C14", FormulaRule(formula=['C14="STOP"'], fill=RED_FILL, font=Font(bold=True, color="9C0006")))

section(wsB, "B17", "KEY METRICS (all closed trades)", 3)
put(wsB, "B18", "Metric", bold=True)
put(wsB, "C18", "Value", bold=True)
put(wsB, "D18", "Formula / meaning", bold=True)
kpis = [
    ("Total closed trades", "=COUNT(TL_Net)", "0", "N = COUNT(Net P&L)"),
    ("Wins", '=COUNTIF(TL_Net,">0")', "0", "Trades with Net P&L > 0"),
    ("Losses", '=COUNTIF(TL_Net,"<0")', "0", "Trades with Net P&L < 0"),
    ("Win rate", "=IFERROR(C20/C19,0)", PCT, "Wins ÷ Total trades"),
    ("Average win ₹", '=IFERROR(AVERAGEIF(TL_Net,">0"),0)', INR, "Σ winning Net ÷ Wins"),
    ("Average loss ₹", '=IFERROR(AVERAGEIF(TL_Net,"<0"),0)', INR, "Σ losing Net ÷ Losses"),
    ("Payoff ratio (avg win / avg loss)", "=IFERROR(C23/ABS(C24),0)", "0.00", "Avg win ÷ |Avg loss|. Sellers usually < 1, so win rate must be high."),
    ("Expectancy ₹ per trade", "=IFERROR(C22*C23-(C21/C19)*ABS(C24),0)", INR, "Win% × Avg win − Loss% × |Avg loss|"),
    ("Expectancy in R (average R-multiple)", "=IFERROR(AVERAGE(TL_R),0)", RFMT, "Mean of R-multiples (Van Tharp). > 0 = edge."),
    ("Profit factor", '=IFERROR(SUMIF(TL_Net,">0")/ABS(SUMIF(TL_Net,"<0")),0)', "0.00", "Gross wins ÷ |Gross losses|. Target > 1.5"),
    ("Net P&L ₹", "=SUM(TL_Net)", INR, "Σ Net P&L (after charges)"),
    ("Total charges ₹", "=SUM(TL_Chg)", INR, "Σ brokerage + STT + exchange + SEBI + stamp + GST"),
    ("Charges as % of gross P&L (abs)", "=IFERROR(C30/ABS(SUM(TL_Gross)),0)", PCT, "Friction drag. High % = overtrading."),
    ("Current equity ₹", "=Capital+C29", INR, "Capital + Net P&L"),
    ("Max drawdown ₹", "=MAX(0,MAX(TL_DD))", INR, "Max of (Peak equity − Equity), trade by trade"),
    ("Max drawdown %", "=MAX(0,MAX(TL_DDPct))", PCT, "Max DD ₹ ÷ Peak equity at that point"),
    ("Recovery factor", "=IFERROR(C29/C33,0)", "0.00", "Net P&L ÷ Max DD. > 2 is healthy; < 0 = still underwater."),
    ("Largest loss ₹", "=MIN(0,MIN(TL_Net))", INR, "Most negative single-trade Net P&L"),
    ("Largest loss as % of capital", "=C36/Capital", PCT, "Rule: never worse than −(Risk%) … your old −24% day."),
    ("Largest win ₹", "=MAX(0,MAX(TL_Net))", INR, ""),
    ("Max consecutive losses", "=MAX(0,MAX(TL_Streak))", "0", "Longest run of losing trades"),
    ("P&L of top 5 trades ₹", "=IFERROR(LARGE(TL_Net,1),0)+IFERROR(LARGE(TL_Net,2),0)+IFERROR(LARGE(TL_Net,3),0)+IFERROR(LARGE(TL_Net,4),0)+IFERROR(LARGE(TL_Net,5),0)", INR, "Σ five largest Net P&L"),
    ("% of Net P&L from top 5 trades", '=IF(C29<=0,"n/a (net ≤ 0)",C40/C29)', PCT, "Top-5 ÷ Net. > 100% = the rest of your trades lose money."),
    ("Total profit given back ₹", "=SUM(TL_GiveBack)", INR, "Σ max(0, MFE ₹ − Gross P&L)"),
    ("Given back as % of MFE", "=IFERROR(C42/SUM(TL_MFE),0)", PCT, "Share of open profit you returned"),
    ("Rule adherence %", '=IFERROR(COUNTIF(TL_Followed,"Y")/(COUNTIF(TL_Followed,"Y")+COUNTIF(TL_Followed,"N")),0)', PCT, "Y ÷ (Y + N). Target 100%."),
    ("Net P&L – rules followed ₹", '=SUMIFS(TL_Net,TL_Followed,"Y")', INR, "What your system actually earns"),
    ("Net P&L – rules broken ₹ (cost of indiscipline)", '=SUMIFS(TL_Net,TL_Followed,"N")', INR, "What breaking rules cost you"),
    ("Oversized / no-SL trades", '=COUNTIF(TL_SizeChk,"OVERSIZED")+COUNTIF(TL_SizeChk,"NO SL")', "0", "From Size Check column"),
    ("Expiry-cutoff breaches", '=COUNTIF(TL_ExpBreach,"BREACH")', "0", "Shorts held past cutoff on DTE 0"),
    ("Days with a limit broken", '=COUNTIF(DS_Status,"LIMIT BROKEN")', "0", "From DAILY SUMMARY"),
]
for i, (lab, f, fmt, note) in enumerate(kpis):
    r = 19 + i
    put(wsB, f"B{r}", lab)
    put(wsB, f"C{r}", f, fmt)
    put(wsB, f"D{r}", note)
KPI_LAST = 19 + len(kpis) - 1
assert wsB["B29"].value == "Net P&L ₹" and wsB["B33"].value == "Max drawdown ₹" and wsB["B40"].value.startswith("P&L of top 5")
for cell in ("C29", "C36", "C46", "C27", "C35"):
    wsB.conditional_formatting.add(cell, CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
wsB.conditional_formatting.add("C37", FormulaRule(formula=["C37<-RiskPct"], fill=RED_FILL))
wsB.conditional_formatting.add("C28", FormulaRule(formula=["AND(C19>0,C28<1)"], fill=RED_FILL))
wsB.conditional_formatting.add("C44", FormulaRule(formula=["C44<1"], fill=AMBER_FILL))

# --- breakdown tables
BR_COLS = ["Trades", "Win %", "Net P&L ₹", "Avg R", "Given Back ₹"]


def breakdown(top, heading, rng_name, items):
    """items: list of (label_formula_or_value, criteria_is_label). Returns last row."""
    section(wsB, f"F{top}", heading, 6)
    put(wsB, f"F{top + 1}", "Category", bold=True)
    for j, h in enumerate(BR_COLS):
        put(wsB, f"{get_column_letter(7 + j)}{top + 1}", h, bold=True)
    r = top + 2
    for lab in items:
        put(wsB, f"F{r}", lab)
        crit = f"$F{r}"
        put(wsB, f"G{r}", f'=IF({crit}="","",COUNTIFS({rng_name},{crit},TL_Net,">-1E+15"))', "0")
        put(wsB, f"H{r}", f'=IF({crit}="","",IFERROR(COUNTIFS({rng_name},{crit},TL_Net,">0")/G{r},0))', PCT)
        put(wsB, f"I{r}", f'=IF({crit}="","",SUMIFS(TL_Net,{rng_name},{crit}))', INR)
        put(wsB, f"J{r}", f'=IF({crit}="","",IFERROR(AVERAGEIFS(TL_R,{rng_name},{crit}),0))', RFMT)
        put(wsB, f"K{r}", f'=IF({crit}="","",SUMIFS(TL_GiveBack,{rng_name},{crit}))', INR)
        r += 1
    rng_ = f"I{top + 2}:I{r - 1}"
    wsB.conditional_formatting.add(rng_, CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    wsB.conditional_formatting.add(rng_, CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL))
    return r - 1


def list_items(list_name, n):
    return [f'=IFERROR(INDEX({list_name},{i}),"")' for i in range(1, n + 1)]


b = 6
blocks = {}
b_end = breakdown(b, "P&L BY STRATEGY", "TL_Strat", list_items("List_Strategy", len(LISTS["List_Strategy"][1])))
blocks["strategy"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "P&L BY DAY OF WEEK", "TL_Day", list_items("List_Weekday", 5))
blocks["dow"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "P&L BY DTE (0 = expiry day)", "TL_DTE", [0, 1, 2, 3, 4, ">=5"])
blocks["dte"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "P&L BY ENTRY-TIME BUCKET", "TL_Bucket", list_items("List_BucketLabel", len(buckets)))
blocks["time"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "P&L BY EXIT REASON", "TL_XReason", list_items("List_ExitReason", len(LISTS["List_ExitReason"][1])))
blocks["exit"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "RULES FOLLOWED (Y) vs BROKEN (N)", "TL_Followed", ["Y", "N"])
blocks["rules"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "COST BY RULE BROKEN", "TL_Broken", list_items("List_RuleBroken", len(LISTS["List_RuleBroken"][1])))
blocks["broken"] = (b + 2, b_end)
b = b_end + 2
b_end = breakdown(b, "P&L BY SETUP GRADE", "TL_Grade", list_items("List_Grade", 3))
blocks["grade"] = (b + 2, b_end)
DASH_LAST = b_end

# --- chart data (rolling window of last N sessions; padding rows show starting capital)
WIN = 40
CD = "N"  # chart-data block starts at column N
wsB.column_dimensions["M"].width = 2
section(wsB, f"{CD}6", f"CHART DATA – last {WIN} sessions (automatic)", 4)
put(wsB, "N7", "Sessions logged", bold=True)
put(wsB, "O7", "=COUNT(DS_Date)", "0")
for j, h in enumerate(["Idx", "Session", "Equity ₹", "Drawdown %"]):
    put(wsB, f"{get_column_letter(14 + j)}8", h, bold=True)
for i in range(1, WIN + 1):
    r = 8 + i
    put(wsB, f"N{r}", f"=$O$7-{WIN}+{i}", "0")
    put(wsB, f"O{r}", f'=IF(N{r}<1,"",TEXT(INDEX(DS_Date,N{r}),"dd-mmm"))')
    put(wsB, f"P{r}", f"=IF(N{r}<1,Capital,INDEX(DS_Eq,N{r}))", INR)
    put(wsB, f"Q{r}", f"=IF(N{r}<1,0,-INDEX(DS_DDPct,N{r}))", PCT)
for c_ in "NOPQ":
    wsB.column_dimensions[c_].width = 11
CD_FIRST, CD_LAST = 9, 8 + WIN

# --- charts (placed below KPI block, left side)
eq = LineChart()
eq.title = f"Equity curve (end of day, last {WIN} sessions)"
eq.y_axis.title = "₹"
eq.y_axis.numFmt = "#,##0"
eq.height, eq.width = 7.5, 17
eq.add_data(Reference(wsB, min_col=16, min_row=8, max_row=CD_LAST), titles_from_data=True)
eq.set_categories(Reference(wsB, min_col=15, min_row=CD_FIRST, max_row=CD_LAST))
eq.legend = None
wsB.add_chart(eq, f"B{KPI_LAST + 3}")

ddc = BarChart()
ddc.title = "Drawdown from peak (%)"
ddc.y_axis.numFmt = "0%"
ddc.height, ddc.width = 6.5, 17
ddc.add_data(Reference(wsB, min_col=17, min_row=8, max_row=CD_LAST), titles_from_data=True)
ddc.set_categories(Reference(wsB, min_col=15, min_row=CD_FIRST, max_row=CD_LAST))
ddc.legend = None
ddc.gapWidth = 20
ddc.series[0].graphicalProperties.solidFill = "C00000"
wsB.add_chart(ddc, f"B{KPI_LAST + 19}")

for key, ttl, anchor in (("exit", "Net P&L by exit reason", f"B{KPI_LAST + 33}"),
                         ("dte", "Net P&L by DTE", f"B{KPI_LAST + 48}"),
                         ("rules", "Rules followed (Y) vs broken (N)", f"B{KPI_LAST + 63}")):
    a, z = blocks[key]
    bc = BarChart()
    bc.type = "bar" if key == "exit" else "col"
    bc.title = ttl
    bc.height, bc.width = 7, 17
    bc.add_data(Reference(wsB, min_col=9, min_row=a, max_row=z), titles_from_data=False)
    bc.set_categories(Reference(wsB, min_col=6, min_row=a, max_row=z))
    bc.legend = None
    bc.y_axis.numFmt = "#,##0"
    bc.series[0].graphicalProperties.solidFill = "2F5597"
    wsB.add_chart(bc, anchor)

for ch in wsB._charts:  # show axes in Excel and keep category labels clear of negative bars
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblPos = "low"

# ======================================================================
# WEEKLY REVIEW
# ======================================================================
wsW = wb.create_sheet("WEEKLY REVIEW", 5)
wsW.sheet_view.showGridLines = False
title(wsW, "WEEKLY REVIEW – every Friday/weekend (≈20 min)",
      "Change the week-start date (Monday) to review any week. Copy this sheet to keep an archive.")
for k, v in {"A": 2, "B": 40, "C": 16, "D": 30, "E": 16, "F": 40}.items():
    wsW.column_dimensions[k].width = v
put(wsW, "B4", "Week starting (Monday)", bold=True)
put(wsW, "C4", dt.date(2026, 9, 28), DATE, inp=True)
put(wsW, "D4", '=IF(C4="","","to "&TEXT(C4+4,"dd-mmm-yy")&IF(WEEKDAY(C4,3)<>0,"  ⚠ not a Monday",""))')
name("WR_Week", "'WEEKLY REVIEW'!$C$4")

section(wsW, "B6", "AUTO STATS FOR THE WEEK", 2)
wk = 'TL_Week,WR_Week'
wstats = [
    (7, "Trades", f'=COUNTIFS({wk},TL_Net,">-1E+15")', "0"),
    (8, "Win rate", f'=IFERROR(COUNTIFS({wk},TL_Net,">0")/C7,0)', PCT),
    (9, "Net P&L ₹", f"=SUMIFS(TL_Net,{wk})", INR),
    (10, "Charges ₹", f"=SUMIFS(TL_Chg,{wk})", INR),
    (11, "Average R", f"=IFERROR(AVERAGEIFS(TL_R,{wk}),0)", RFMT),
    (12, "Rule adherence %", f'=IFERROR(COUNTIFS({wk},TL_Followed,"Y")/(COUNTIFS({wk},TL_Followed,"Y")+COUNTIFS({wk},TL_Followed,"N")),0)', PCT),
    (13, "Rule-break cost ₹ (net P&L of rule-broken trades)", f'=SUMIFS(TL_Net,{wk},TL_Followed,"N")', INR),
    (14, "Profit given back ₹", f"=SUMIFS(TL_GiveBack,{wk})", INR),
    (15, "Days with a limit broken", '=COUNTIFS(DS_Week,WR_Week,DS_Status,"LIMIT BROKEN")', "0"),
    (16, "Week net vs weekly limit", '=IF(C9<=-WeeklyDDLimit,"BREACH – no trading till next week","Within limit")', None),
]
for r, lab, f, fmt in wstats:
    put(wsW, f"B{r}", lab)
    put(wsW, f"C{r}", f, fmt)
for cell in ("C9", "C13"):
    wsW.conditional_formatting.add(cell, CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
wsW.conditional_formatting.add("C16", FormulaRule(formula=['LEFT(C16,6)="BREACH"'], fill=RED_FILL))

section(wsW, "E6", "BEST & WORST TRADE (auto)", 2)
best_worst = [
    ("Best", "_xlfn.MAXIFS"), ("Worst", "_xlfn.MINIFS")]
r0 = 7
for lab, fn in best_worst:
    pos = f"IFERROR(MATCH(1,INDEX((TL_Week=WR_Week)*(TL_Net={fn}(TL_Net,{wk})),0),0),0)"
    put(wsW, f"E{r0}", f"{lab} trade Net ₹", bold=True)
    put(wsW, f"F{r0}", f'=IF(C7=0,"",{fn}(TL_Net,{wk}))', INR)
    put(wsW, f"E{r0 + 1}", f"{lab} trade #")
    put(wsW, f"F{r0 + 1}", f'=IF(C7=0,"",INDEX(TL_No,{pos}))', "0")
    put(wsW, f"E{r0 + 2}", "Details")
    put(wsW, f"F{r0 + 2}",
        f'=IF(C7=0,"",TEXT(INDEX(TL_Date,{pos}),"ddd dd-mmm")&" | "&INDEX(TL_Strike,{pos})&" "&INDEX(TL_Type,{pos})&" | "&INDEX(TL_Strat,{pos})&" | exit: "&INDEX(TL_XReason,{pos}))')
    r0 += 4
wsW.conditional_formatting.add("F11", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))

section(wsW, "B18", "RULE-BREAK COST BY RULE (auto) – use this to pick your top 3 mistakes", 3)
put(wsW, "B19", "Rule broken", bold=True)
put(wsW, "C19", "Times", bold=True)
put(wsW, "D19", "Net P&L of those trades ₹", bold=True)
nrules = len(LISTS["List_RuleBroken"][1])
for i in range(2, nrules + 1):  # skip "None"
    r = 18 + i
    put(wsW, f"B{r}", f'=IFERROR(INDEX(List_RuleBroken,{i}),"")')
    put(wsW, f"C{r}", f'=IF(B{r}="","",COUNTIFS({wk},TL_Broken,B{r}))', "0")
    put(wsW, f"D{r}", f'=IF(B{r}="","",SUMIFS(TL_Net,{wk},TL_Broken,B{r}))', INR)
rb_last = 18 + nrules
wsW.conditional_formatting.add(f"D20:D{rb_last}", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))

r = rb_last + 2
section(wsW, f"B{r}", "TOP 3 MISTAKES (write them)", 5)
put(wsW, f"B{r + 1}", "Mistake", bold=True)
put(wsW, f"C{r + 1}", "Cost ₹", bold=True)
put(wsW, f"D{r + 1}", "Rule #", bold=True)
put(wsW, f"E{r + 1}", "Trigger (what I felt)", bold=True)
put(wsW, f"F{r + 1}", "Fix / if-then plan", bold=True)
mistakes = [
    ("Held 25850 CE spread past 14:45 on expiry", -6175, "R4 Held short past expiry cutoff", "Greedy",
     "IF it is 14:40 on expiry THEN square off all shorts, no exceptions."),
    ("Did not trail SL after 51 → 30 (profit given back)", -13000, "R8 Gave back profit (no trail)", "Greedy",
     "IF premium falls 40% THEN move SL to cost; at 60% lock half."),
    ("5 lots vs cap of 2", -3705, "R1 Oversized position", "Overconfidence",
     "IF lots > 'Lots allowed' on RULES THEN cancel the order."),
]
for i, (m, cost, rule, trig, fix) in enumerate(mistakes):
    rr = r + 2 + i
    put(wsW, f"B{rr}", m, inp=True)
    put(wsW, f"C{rr}", cost, INR, inp=True)
    put(wsW, f"D{rr}", rule, inp=True)
    put(wsW, f"E{rr}", trig, inp=True)
    put(wsW, f"F{rr}", fix, inp=True)
dv = DataValidation(type="list", formula1="=List_RuleBroken", allow_blank=True)
wsW.add_data_validation(dv)
dv.add(f"D{r + 2}:D{r + 4}")

r = r + 6
section(wsW, f"B{r}", "REFLECTION", 5)
refl = [
    ("Best trade – why it worked", "Followed the plan: spread, SL order placed, booked at target."),
    ("Worst trade – what went wrong", "Expiry-day short held into the last hour; mental SL, 5 lots."),
    ("Rule-break cost ₹ (auto, from above)", f"=C13"),
    ("ONE change for next week (specific, measurable)", "Exit every expiry-day short by 14:45 via a GTT/alarm at 14:40."),
    ("Did I keep last week's ONE change? (Y/N)", "N"),
    ("Discipline score this week (1-5)", 2),
    ("Next week's max lots (may reduce after a losing week)", "=MaxLots"),
]
for i, (lab, v) in enumerate(refl):
    rr = r + 1 + i
    put(wsW, f"B{rr}", lab, bold=True)
    is_f = isinstance(v, str) and v.startswith("=")
    put(wsW, f"C{rr}", v, INR if lab.startswith("Rule-break") else None, inp=not is_f)
    wsW.merge_cells(f"C{rr}:F{rr}")
yn_row = r + 5
dv2 = DataValidation(type="list", formula1="=List_YN", allow_blank=True)
wsW.add_data_validation(dv2)
dv2.add(f"C{yn_row}")
dv3 = DataValidation(type="list", formula1="=List_Scale", allow_blank=True)
wsW.add_data_validation(dv3)
dv3.add(f"C{yn_row + 1}")

# ======================================================================
# final touches
# ======================================================================
order = ["RULES", "PRE-MARKET", "TRADE LOG", "DAILY SUMMARY", "DASHBOARD", "WEEKLY REVIEW", "LISTS"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
tab_colors = {"RULES": "1F3864", "PRE-MARKET": "2F5597", "TRADE LOG": "2F5597", "DAILY SUMMARY": "404040",
              "DASHBOARD": "C00000", "WEEKLY REVIEW": "548235", "LISTS": "A6A6A6"}
for n, c in tab_colors.items():
    wb[n].sheet_properties.tabColor = c
for ws in wb.worksheets:
    for row in ws.iter_rows(min_row=1, max_row=2):
        for c in row:
            if c.value is not None and c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size, bold=c.font.bold, color=c.font.color)
    ws.sheet_view.zoomScale = 90
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
